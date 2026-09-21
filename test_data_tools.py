"""Regression tests for reviewed updates and lossless maintenance operations."""
from copy import deepcopy
import hashlib
import json
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest

import openpyxl
from apply_changes import apply_changes
from data_tools import ROOT, atomic_write, read_json
from diff_workbooks import compare_workbooks
from export_xlsx import build_workbook
from format_json import format_data
from validate_data import validate_data
from xlsx_tools import load_workbook_data


def fixture():
    return {'_version': 1, '_comment': 'quote " and \\ and\nnewline', 'tier_info': {},
            'extra': {'keep': True}, 'tracks': [
                {'大地图': big, '小地图': 'Same', 'note': 'keep " me',
                 zone: {tier: [] for tier in ('理论', '高手', '普通', '自动')}}
                for big, zone in [('A', '五区'), ('B', '五区')]]}


CARS = {'cars': [{'title': name} for name in ('X', 'Y', 'Z', '恶魔')]}


def sc_fixture():
    """两区齐全的赛道骨架, 供特殊跑法透视表测试使用"""
    return {'_version': 1, 'tracks': [
        {'大地图': big, '小地图': 'Same',
         **{zone: {tier: [] for tier in ('理论', '高手', '普通', '自动')} for zone in ('五区', '四区')}}
        for big in ('A', 'B')]}


def change(**overrides):
    value = {'accepted': True, 'big': 'A', 'small': 'Same', 'zone': '五区',
             'tier': '高手', 'car': 'X', 'stars': 6, 'sc': False, 'sc_type': None,
             'old_present': False, 'new_present': True, 'old': None, 'new': 20}
    value.update(overrides)
    return value


def report(*changes):
    return {'version': 1, 'changes': list(changes), 'structure': []}


class DataToolsTests(unittest.TestCase):
    def test_format_preserves_unknown_fields_escaping_and_missing_zones(self):
        data = fixture()
        self.assertEqual(json.loads(format_data(data)), data)
        self.assertEqual(format_data(json.loads(format_data(data))), format_data(data))
        with self.assertRaises(ValueError):
            format_data({'time': float('nan')})

    def test_real_data_validation_and_format(self):
        data = read_json(ROOT / 'gauntlet_data.json')
        errors, counts = validate_data(data, read_json(ROOT / 'cars.json'))
        self.assertEqual(errors, [])
        self.assertGreater(sum(counts.values()), 0)
        self.assertEqual(json.loads(format_data(data)), data)

    def test_rejected_changes_and_no_source_mutation(self):
        data = fixture()
        original = deepcopy(data)
        result, logs = apply_changes(data, report(change(accepted=False)), CARS)
        self.assertEqual(result, original)
        self.assertEqual(logs, [])
        apply_changes(data, report(change()), CARS)
        self.assertEqual(data, original)

    def test_placeholder_fill_mirror_and_composite_track_key(self):
        data = fixture()
        data['tracks'][0]['五区']['高手'] = [{'cars': [{'name': 'X', 'stars': 6}], 'note': 'retain'}]
        result, _ = apply_changes(data, report(change()), CARS)
        tiers = result['tracks'][0]['五区']
        self.assertEqual(len(tiers['高手']), 1)
        self.assertEqual(tiers['高手'][0]['note'], 'retain')
        self.assertEqual(tiers['高手'][0]['time'], 20)
        self.assertEqual(len(tiers['普通']), 1)
        self.assertEqual(result['tracks'][1]['五区']['高手'], [])

    def test_theory_and_sc_do_not_add_mirrors(self):
        for value in (change(tier='理论'), change(sc=True, sc_type='跳图')):
            result, _ = apply_changes(fixture(), report(value), CARS)
            self.assertEqual(result['tracks'][0]['五区']['普通'], [])

    def test_sc_placeholder_fill_accepts_absent_old_and_keeps_route(self):
        """sc 占位条目(有 sc_type 无 time)对应 xlsx 空单元格, 填值应视作新增而非旧值不匹配"""
        data = fixture()
        data['tracks'][0]['五区']['理论'] = [{'cars': [{'name': 'X'}], 'sc': True, 'sc_type': '跳图'}]
        result, _ = apply_changes(data, report(
            change(tier='理论', stars=None, sc=True, sc_type='跳图', new=15)), CARS)
        entries = result['tracks'][0]['五区']['理论']
        self.assertEqual(len(entries), 1)
        self.assertEqual(entries[0]['sc_type'], '跳图')
        self.assertEqual(entries[0]['time'], 15)
        self.assertEqual(result['tracks'][0]['五区']['普通'], [])

    def test_deletion_removes_expert_mirror_and_slower_update_is_explicit(self):
        data, _ = apply_changes(fixture(), report(change()), CARS)
        slower = change(old_present=True, old=20, new=21)
        updated, _ = apply_changes(data, report(slower), CARS)
        self.assertEqual(updated['tracks'][0]['五区']['高手'][0]['time'], 21)
        deleted, _ = apply_changes(data, report(change(old_present=True, old=20, new=None, new_present=False)), CARS)
        self.assertEqual(deleted['tracks'][0]['五区']['高手'], [])
        self.assertEqual(deleted['tracks'][0]['五区']['普通'], [])

    def test_conflicts_drift_unknown_car_invalid_time_and_stars_fail(self):
        data, _ = apply_changes(fixture(), report(change()), CARS)
        with self.assertRaisesRegex(ValueError, '旧值不匹配'):
            apply_changes(data, report(change(old_present=True, old=19)), CARS)
        with self.assertRaisesRegex(ValueError, '多个改动'):
            apply_changes(fixture(), report(change(), change(new=21)), CARS)
        for value in (change(car='missing'), change(stars=7), change(new=True), change(new=-1), change(accepted='true')):
            with self.assertRaises(ValueError):
                apply_changes(fixture(), report(value), CARS)

    def test_sorting_and_special_route_identity(self):
        result, _ = apply_changes(fixture(), report(change(new=30), change(car='Y', new=20),
                                                   change(tier='理论', sc=True, sc_type='跳图', new=10),
                                                   change(tier='理论', sc=True, sc_type='滑雪', new=11)), CARS)
        self.assertEqual([e['time'] for e in result['tracks'][0]['五区']['高手']], [20, 30])
        result, _ = apply_changes(result, report(change(tier='理论', sc=True, sc_type='跳图', old_present=True, old=10, new=9)), CARS)
        self.assertEqual([e['time'] for e in result['tracks'][0]['五区']['理论']], [9, 11])

    def test_special_route_pivot_sheets_share_rows_with_per_tier_columns(self):
        data, _ = apply_changes(sc_fixture(), report(
            change(tier='理论', car='X', stars=None, sc=True, sc_type='跳图', new=20),
            change(tier='理论', car='Y', stars=None, sc=True, sc_type='跳图', new=21),
            change(zone='四区', tier='高手', sc=True, sc_type='跳图', new=19),
            change(big='B', sc=True, sc_type='滑雪', new=30)), CARS)
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / 'data.xlsx'
            workbook, _ = build_workbook(data['tracks'])
            workbook.save(path)
            workbook.close()
            sheets = load_workbook_data(path)
        for zone in ('五区', '四区'):
            for tier in ('理论', '高手'):
                name = f'{zone}_{tier}_特殊跑法'
                self.assertIn(name, sheets)
                # 全部跑法组合在四张表都出现, 空行可直接填值
                self.assertEqual(set(sheets[name]['rows']), {('A', 'Same', '跳图'), ('B', 'Same', '滑雪')}, name)
        # 列集按区档各自收集: 理论档无星列, 高手档带星列
        self.assertEqual(sheets['五区_理论_特殊跑法']['headers'], ['X', 'Y'])
        self.assertEqual(sheets['五区_高手_特殊跑法']['headers'], ['X★6'])
        self.assertEqual(sheets['四区_理论_特殊跑法']['headers'], [])
        self.assertEqual(sheets['四区_高手_特殊跑法']['headers'], ['X★6'])
        # 有值格只落在本区本档的列上
        self.assertEqual(sheets['五区_理论_特殊跑法']['rows'][('A', 'Same', '跳图')], {'X': 20, 'Y': 21})
        self.assertEqual(sheets['五区_理论_特殊跑法']['rows'][('B', 'Same', '滑雪')], {})
        self.assertEqual(sheets['五区_高手_特殊跑法']['rows'][('B', 'Same', '滑雪')], {'X★6': 30})
        self.assertEqual(sheets['五区_高手_特殊跑法']['rows'][('A', 'Same', '跳图')], {})
        self.assertEqual(sheets['四区_高手_特殊跑法']['rows'][('A', 'Same', '跳图')], {'X★6': 19})
        self.assertEqual(sheets['四区_理论_特殊跑法']['rows'][('A', 'Same', '跳图')], {})
        # 只有特殊跑法成绩的车不占主表列
        self.assertEqual(sheets['五区_理论']['headers'], [])
        self.assertEqual(sheets['五区_高手']['headers'], [])
        self.assertEqual(sheets['四区_理论']['headers'], [])

    def test_special_route_edit_review_apply_round_trip(self):
        data, _ = apply_changes(sc_fixture(), report(
            change(tier='理论', car='X', stars=None, sc=True, sc_type='跳图', new=20),
            change(tier='理论', car='Y', stars=None, sc=True, sc_type='跳图', new=21)), CARS)
        with tempfile.TemporaryDirectory() as temp:
            base, user = [Path(temp) / name for name in ('base.xlsx', 'user.xlsx')]
            workbook, _ = build_workbook(data['tracks'])
            self.assertEqual(set(workbook.sheetnames), {f'{zone}_{tier}' for zone in ('五区', '四区')
                             for tier in ('理论', '高手', '普通', '自动')} | {f'{zone}_{tier}_特殊跑法'
                             for zone in ('五区', '四区') for tier in ('理论', '高手')})
            workbook.save(base)
            self.assertEqual(compare_workbooks(load_workbook_data(base), load_workbook_data(base))['changes'], [])
            sheet = workbook['五区_理论_特殊跑法']
            # 理论档列集无星, 两车按成绩升序排
            self.assertEqual(sheet.cell(1, 4).value, 'X')
            self.assertEqual(sheet.cell(1, 5).value, 'Y')
            self.assertEqual(sheet.cell(2, 4).value, 20)
            sheet.cell(2, 4).value = 19.5      # 改值
            sheet.cell(2, 5).value = 22        # 改值
            workbook.save(user)
            workbook.close()
            changes = compare_workbooks(load_workbook_data(base), load_workbook_data(user))
            self.assertEqual(len(changes['changes']), 2)
            for finding in changes['changes']:
                self.assertTrue(finding['sc'])
                self.assertEqual(finding['sc_type'], '跳图')
                self.assertEqual((finding['zone'], finding['tier']), ('五区', '理论'))
                self.assertIsNone(finding['stars'])
                finding['accepted'] = True
            result, _ = apply_changes(data, changes, CARS)
            times = {e['cars'][0]['name']: e['time'] for e in result['tracks'][0]['五区']['理论']}
            self.assertEqual(times, {'X': 19.5, 'Y': 22})
            # 删值: 清空 Y 的单元格 → 删除条目
            fresh, _ = build_workbook(result['tracks'])
            fresh['五区_理论_特殊跑法'].cell(2, 5).value = None
            fresh.save(base)
            fresh.close()
            deletion = compare_workbooks(load_workbook_data(user), load_workbook_data(base))
            self.assertEqual([c['kind'] for c in deletion['changes']], ['deleted'])
            deletion['changes'][0]['accepted'] = True
            trimmed, _ = apply_changes(result, deletion, CARS)
            self.assertEqual(len(trimmed['tracks'][0]['五区']['理论']), 1)

    def test_special_route_high_tier_keeps_starred_columns(self):
        """高手档带星条目进高手表带星列, 且不混入理论表"""
        data, _ = apply_changes(sc_fixture(), report(
            change(tier='高手', car='X', stars=6, sc=True, sc_type='跳图', new=20),
            change(tier='理论', car='Y', stars=None, sc=True, sc_type='跳图', new=21)), CARS)
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / 'data.xlsx'
            workbook, _ = build_workbook(data['tracks'])
            workbook.save(path)
            workbook.close()
            sheets = load_workbook_data(path)
        self.assertEqual(sheets['五区_理论_特殊跑法']['headers'], ['Y'])
        self.assertEqual(sheets['五区_高手_特殊跑法']['headers'], ['X★6'])
        self.assertEqual(sheets['五区_理论_特殊跑法']['rows'][('A', 'Same', '跳图')], {'Y': 21})
        self.assertEqual(sheets['五区_高手_特殊跑法']['rows'][('A', 'Same', '跳图')], {'X★6': 20})

    def test_validation_requires_known_special_route_type(self):
        data, _ = apply_changes(sc_fixture(), report(change(tier='理论', sc=True, sc_type='跳图', new=20)), CARS)
        entry = data['tracks'][0]['五区']['理论'][0]
        self.assertEqual(validate_data(data, CARS)[0], [])
        for sc_type, message in ((None, '缺类型'), ('jump', '未知')):
            entry['sc_type'] = sc_type
            errors = validate_data(data, CARS)[0]
            self.assertTrue(errors, sc_type)
            self.assertIn(message, errors[0])

    def test_expert_and_explicit_mirror_changes_share_original_snapshot(self):
        mirror = change(tier='普通', new='✓')
        result, _ = apply_changes(fixture(), report(change(), mirror), CARS)
        self.assertEqual(len(result['tracks'][0]['五区']['普通']), 1)
        deletion = change(old_present=True, old=20, new_present=False, new=None)
        mirror_deletion = change(tier='普通', old_present=True, old='✓', new_present=False, new=None)
        result, _ = apply_changes(result, report(deletion, mirror_deletion), CARS)
        self.assertEqual(result['tracks'][0]['五区']['普通'], [])

    def test_diff_review_apply_export_end_to_end(self):
        data, _ = apply_changes(fixture(), report(change(tier='理论')), CARS)
        with tempfile.TemporaryDirectory() as temp:
            base, user = [Path(temp) / name for name in ('base.xlsx', 'user.xlsx')]
            workbook, _ = build_workbook(data['tracks'])
            workbook.save(base)
            workbook['五区_理论'].cell(2, 3).value = 19
            workbook.save(user)
            workbook.close()
            changes = compare_workbooks(load_workbook_data(base), load_workbook_data(user))
            self.assertEqual(len(changes['changes']), 1)
            self.assertEqual(changes['changes'][0]['kind'], 'faster')
            self.assertFalse(changes['changes'][0]['accepted'])
            changes['changes'][0]['accepted'] = True
            result, _ = apply_changes(data, changes, CARS)
            workbook, _ = build_workbook(result['tracks'])
            workbook.save(base)
            workbook.close()
            self.assertEqual(compare_workbooks(load_workbook_data(base), load_workbook_data(user))['changes'], [])

    def test_columns_can_move_and_missing_columns_are_not_deletions(self):
        data, _ = apply_changes(fixture(), report(change(), change(car='Y', new=21)), CARS)
        with tempfile.TemporaryDirectory() as temp:
            base, user = [Path(temp) / name for name in ('base.xlsx', 'user.xlsx')]
            workbook, _ = build_workbook(data['tracks'])
            workbook.save(base)
            sheet = workbook['五区_高手']
            for row in range(1, sheet.max_row + 1):
                left, right = sheet.cell(row, 3).value, sheet.cell(row, 4).value
                sheet.cell(row, 3).value = right
                sheet.cell(row, 4).value = left
            workbook.save(user)
            self.assertEqual(compare_workbooks(load_workbook_data(base), load_workbook_data(user))['changes'], [])
            sheet.delete_cols(4)
            workbook.save(user)
            workbook.close()
            changes = compare_workbooks(load_workbook_data(base), load_workbook_data(user))
            self.assertEqual(changes['changes'], [])
            self.assertEqual(changes['structure'][0]['kind'], 'column_missing')

    def test_validation_detects_each_business_error(self):
        data, _ = apply_changes(fixture(), report(change()), CARS)
        for mode in ('duplicate', 'star', 'name', 'mirror', 'order', 'time'):
            invalid = deepcopy(data)
            tiers = invalid['tracks'][0]['五区']
            if mode == 'duplicate':
                tiers['高手'].append(deepcopy(tiers['高手'][0]))
            elif mode == 'star':
                tiers['高手'][0]['cars'][0]['stars'] = 0
            elif mode == 'name':
                tiers['高手'][0]['cars'][0]['name'] = 'missing'
            elif mode == 'mirror':
                tiers['普通'].clear()
            elif mode == 'order':
                tiers['理论'] = [{'cars': [{'name': 'X'}], 'time': 30}, {'cars': [{'name': 'Y'}], 'time': 20}]
            else:
                tiers['高手'][0]['time'] = float('inf')
            self.assertTrue(validate_data(invalid, CARS)[0], mode)

    def test_atomic_backup_and_write_failure(self):
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / 'data.json'
            path.write_text('original', encoding='utf-8')
            atomic_write(path, 'replacement')
            self.assertEqual(path.read_text(), 'replacement')
            self.assertEqual(next(Path(temp).glob('*.bak')).read_text(), 'original')
            with self.assertRaises(OSError):
                atomic_write(Path(temp) / 'missing' / 'file.json', 'bad')
            self.assertEqual(path.read_text(), 'replacement')

    def test_workbook_export_and_keyed_reordering(self):
        data, _ = apply_changes(fixture(), report(change()), CARS)
        with tempfile.TemporaryDirectory() as temp:
            left, right = (Path(temp) / n for n in ('left.xlsx', 'right.xlsx'))
            workbook, _ = build_workbook(data['tracks'])
            workbook.save(left)
            # Swap track rows with the same small-map name; keys must preserve the big map.
            sheet = workbook['五区_高手']
            rows = [[c.value for c in sheet[row]] for row in (2, 3)]
            for r, values in zip((2, 3), reversed(rows)):
                for c, value in enumerate(values, 1):
                    sheet.cell(r, c).value = value
            workbook.save(right)
            workbook.close()
            self.assertEqual(compare_workbooks(load_workbook_data(left), load_workbook_data(right))['changes'], [])
            workbook = openpyxl.load_workbook(right)
            workbook.remove(workbook['五区_高手'])
            workbook.save(right)
            workbook.close()
            diff = compare_workbooks(load_workbook_data(left), load_workbook_data(right))
            self.assertEqual(diff['changes'], [])
            self.assertEqual(diff['structure'][0]['kind'], 'sheet_missing')

    def test_duplicate_columns_rows_and_formulas_rejected(self):
        for mode in ('column', 'row', 'formula'):
            with tempfile.TemporaryDirectory() as temp:
                path = Path(temp) / 'bad.xlsx'
                workbook, _ = build_workbook(fixture()['tracks'])
                sheet = workbook['五区_理论']
                if mode == 'column':
                    sheet.cell(1, 3, 'X')
                    sheet.cell(1, 4, 'X')
                elif mode == 'row':
                    sheet.cell(3, 1, 'A')
                else:
                    sheet.cell(2, 3, '=1+1')
                workbook.save(path)
                workbook.close()
                with self.assertRaises(ValueError):
                    load_workbook_data(path)

    def test_cli_review_to_preview_and_protected_inputs(self):
        with tempfile.TemporaryDirectory() as temp:
            directory = Path(temp)
            source, review, cars, preview = [directory / name for name in ('source.json', 'review.json', 'cars.json', 'preview.json')]
            source.write_text(json.dumps(fixture(), ensure_ascii=False), encoding='utf-8')
            cars.write_text(json.dumps(CARS, ensure_ascii=False), encoding='utf-8')
            review.write_text(json.dumps(report(change()), ensure_ascii=False), encoding='utf-8')
            digest = hashlib.sha256(source.read_bytes()).digest()
            command = [sys.executable, str(ROOT / 'apply_changes.py'), '--input', str(source), '--cars', str(cars), '--review', str(review)]
            result = subprocess.run(command, cwd=directory, capture_output=True)
            self.assertEqual(result.returncode, 0, result.stderr)
            self.assertFalse(preview.exists())
            result = subprocess.run(command + ['--output', str(preview)], cwd=directory, capture_output=True)
            self.assertEqual(result.returncode, 0, result.stderr)
            self.assertEqual(read_json(preview)['tracks'][0]['五区']['高手'][0]['time'], 20)
            self.assertEqual(hashlib.sha256(source.read_bytes()).digest(), digest)
            self.assertNotEqual(subprocess.run(command + ['--output', str(source)], capture_output=True).returncode, 0)
            result = subprocess.run(command + ['--write'], capture_output=True)
            self.assertEqual(result.returncode, 0, result.stderr)
            self.assertTrue(list(directory.glob('source.json.*.bak')))


if __name__ == '__main__':
    unittest.main()
