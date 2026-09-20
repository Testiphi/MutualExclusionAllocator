"""Keyed workbook parsing; reject ambiguous rows, columns and formulas."""
import openpyxl
from data_tools import parse_column, parse_sheet_name


def load_workbook_data(path):
    workbook = openpyxl.load_workbook(path, data_only=False)
    result = {}
    try:
        for sheet in workbook:
            for row in sheet:
                if any(cell.data_type == 'f' for cell in row):
                    raise ValueError(f'请将公式转为明确数值: {sheet.title}')
            parsed = parse_sheet_name(sheet.title)
            if parsed is None:
                continue
            zone, tier, sc = parsed
            # 普通表 A大地图 B小地图 C…车辆; 特殊跑法表 C 为跑法行键
            first = 4 if sc else 3
            headers = [sheet.cell(1, c).value for c in range(first, sheet.max_column + 1)]
            present = [h for h in headers if h not in (None, '')]
            if len(set(present)) != len(present):
                raise ValueError(f'重复车辆列: {sheet.title}')
            for header in present:
                parse_column(header)
            rows = {}
            big = None
            for row in sheet.iter_rows(min_row=2, values_only=True):
                if row[0] not in (None, ''):
                    big = row[0]
                if row[1] in (None, ''):
                    continue
                if sc and row[2] in (None, ''):
                    raise ValueError(f'缺跑法类型: {sheet.title} {big}/{row[1]}')
                key = (big, row[1], row[2]) if sc else (big, row[1])
                if big is None or key in rows:
                    raise ValueError(f'缺大地图或重复赛道: {sheet.title} {key}')
                rows[key] = {header: value for header, value in zip(headers, row[first - 1:])
                             if header not in (None, '') and value not in (None, '')}
            result[sheet.title] = {'headers': present, 'rows': rows, 'zone': zone, 'tier': tier, 'sc': sc}
        return result
    finally:
        workbook.close()
